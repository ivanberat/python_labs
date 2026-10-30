import csv
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ENCODING = "utf-8"
SHEET_NAME = "Sheet1"
MIN_WIDTH = 8


def _check_suffix(path: Path, suffix: str) -> None:
    """Проверяет расширение файла, иначе ValueError."""
    if path.suffix.lower() != suffix:
        raise ValueError(f"ожидался файл {suffix}, получен: {path.name}")


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """Конвертирует CSV в XLSX.

    Первая строка CSV — заголовок (жирным). Лист называется Sheet1.
    Ширина колонки — по самой длинной ячейке, но не меньше 8 символов.
    """
    src = Path(csv_path)
    _check_suffix(src, ".csv")
    dst = Path(xlsx_path)
    _check_suffix(dst, ".xlsx")

    with src.open(encoding=ENCODING, newline="") as f:
        # csv.reader отдаёт пустую строку файла как [], такие строки пропускаем
        rows = [row for row in csv.reader(f) if row]

    if not rows:
        raise ValueError("CSV пустой или без заголовка")
    if len(rows) < 2:
        raise ValueError("В CSV нет строк данных, только заголовок")
    width = len(rows[0])
    if any(len(row) != width for row in rows):
        raise ValueError("В CSV строки разной длины, таблица не получится")

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for row in rows:
        ws.append(row)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    # zip(*rows) даёт колонки; ширина — по самой длинной ячейке + запас 2
    for index, column in enumerate(zip(*rows), start=1):
        longest = max(len(value) for value in column)
        ws.column_dimensions[get_column_letter(index)].width = max(MIN_WIDTH, longest + 2)

    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)


def _describe(xlsx_path: Path) -> None:
    """Открывает результат обратно и печатает, что в нём лежит."""
    wb = load_workbook(xlsx_path)
    ws = wb[SHEET_NAME]
    widths = {
        letter: ws.column_dimensions[letter].width
        for letter in (get_column_letter(i) for i in range(1, ws.max_column + 1))
    }
    print(f"  лист: {ws.title}, строк: {ws.max_row}, колонок: {ws.max_column}")
    print(f"  ширины колонок: {widths}")


def _demo() -> None:
    """Сценарий 3: конвертирует примеры из data/samples в data/out и проверяет результат."""
    samples = Path("data") / "samples"
    out = Path("data") / "out"
    for name in ("people", "cities"):
        csv_to_xlsx(samples / f"{name}.csv", out / f"{name}.xlsx")
        with (samples / f"{name}.csv").open(encoding=ENCODING, newline="") as f:
            csv_rows = sum(1 for row in csv.reader(f) if row)
        wb = load_workbook(out / f"{name}.xlsx")
        if wb[SHEET_NAME].max_row != csv_rows:
            raise ValueError(f"число строк в {name}.xlsx не совпадает с CSV")
        print(f"Сценарий 3, CSV → XLSX: {name}.csv → {name}.xlsx")
        _describe(out / f"{name}.xlsx")


def _self_check() -> None:
    """Проверки ошибок на временных файлах. Каждый случай печатает результат."""
    cases = [
        ("пустой CSV-файл", "a.csv", ""),
        ("CSV: только заголовок", "a.csv", "name,age\n"),
        ("CSV: строки разной длины", "a.csv", "a,b\n1,2\n3\n"),
        ("неверное расширение входа", "a.txt", "a\n1\n"),
    ]
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        for title, name, content in cases:
            src = tmp_dir / name
            src.write_text(content, encoding=ENCODING)
            try:
                csv_to_xlsx(src, tmp_dir / "out.xlsx")
                print(f"НЕТ ОШИБКИ: {title}")
            except ValueError as error:
                print(f"ок, {title}: {error}")

        try:
            csv_to_xlsx(tmp_dir / "нет.csv", tmp_dir / "out.xlsx")
        except FileNotFoundError:
            print("ок, файл отсутствует: FileNotFoundError")

        try:
            csv_to_xlsx(tmp_dir / "a.csv", tmp_dir / "out.txt")
        except ValueError as error:
            print(f"ок, неверное расширение выхода: {error}")

        cp = tmp_dir / "cp.csv"
        cp.write_bytes("имя,город\nАлиса,Москва\n".encode("cp1251"))
        try:
            csv_to_xlsx(cp, tmp_dir / "out.xlsx")
        except ValueError as error:
            print(f"ок, не UTF-8: {type(error).__name__}")


def main() -> int:
    try:
        _demo()
    except (OSError, ValueError) as error:
        print(f"Ошибка: {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    _self_check()
    print()
    sys.exit(main())
